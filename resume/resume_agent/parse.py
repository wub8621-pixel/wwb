"""解析层：unstructured 主解析，结果异常时降级到 native。

主备不是“再解析一遍”这么简单，关键是**每次降级都要有可读的理由**。
调用方拿到的 ParseResult 里保留 attempts，能回答“为什么没用主解析”。

native = pymupdf 读 PDF、python-docx 读 DOCX，不依赖版面模型。
"""

from __future__ import annotations

import threading
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from .quality import QualityReport, assess_quality, strip_repeated_edges

# 引擎顺序即优先级：主解析在前，备用在后
ENGINE_ORDER = ("unstructured", "native")

# 单个引擎的时间上限。unstructured 会尝试下载 NLTK 数据，网络不通时能卡几分钟；
# 超时即视为该引擎失败，直接降级，不能让整份简历一直挂着。
ENGINE_TIMEOUTS: dict[str, float] = {"unstructured": 20.0, "native": 120.0}
DEFAULT_TIMEOUT_S = 60.0

Extractor = Callable[[Path], "tuple[list[str], list[list[str]]]"]


class ParseError(RuntimeError):
    """所有解析引擎都不可用，必须显式失败，不能返回空文本来蒙混。"""


@dataclass(frozen=True)
class ParseResult:
    path: Path
    engine: str
    lines: list[str]
    pages: list[list[str]]
    quality: QualityReport
    attempts: tuple[tuple[str, str], ...]

    @property
    def text(self) -> str:
        return "\n".join(line for line in self.lines if line.strip())


def extract_native(path: Path) -> tuple[list[str], list[list[str]]]:
    """pymupdf / python-docx 抽取，保留每页的原始行（含空行）。"""
    if path.suffix.lower() == ".pdf":
        import pymupdf

        pages: list[list[str]] = []
        with pymupdf.open(path) as doc:
            for page in doc:
                pages.append(page.get_text("text").splitlines())
        return [line for page in pages for line in page], pages

    from docx import Document

    document = Document(str(path))
    lines = [p.text for p in document.paragraphs]
    # 简历常用表格排版，单元格内容同样要收进来
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                lines.extend(cell.text.splitlines())
    return lines, [lines]


def extract_unstructured(path: Path) -> tuple[list[str], list[list[str]]]:
    """官方解析器。中文必须显式指定 language，否则会按英文处理并返回空结果。"""
    try:
        from unstructured.partition.auto import partition
    except ImportError as exc:
        raise RuntimeError("未安装 unstructured") from exc

    kwargs: dict[str, object] = {"filename": str(path), "languages": ["chi_sim"]}
    if path.suffix.lower() == ".pdf":
        kwargs["strategy"] = "fast"

    elements = partition(**kwargs)
    by_page: dict[int, list[str]] = {}
    for element in elements:
        page = getattr(element.metadata, "page_number", None) or 1
        by_page.setdefault(int(page), []).append(element.text)
    pages = [by_page[number] for number in sorted(by_page)]
    return [line for page in pages for line in page], pages


EXTRACTORS: dict[str, Extractor] = {
    "unstructured": extract_unstructured,
    "native": extract_native,
}


def call_with_timeout(func: Callable[[Path], "tuple[list[str], list[list[str]]]"], path: Path, timeout: float):
    """在守护线程里执行抽取，超时抛 TimeoutError。

    解析器可能阻塞在网络 IO 上（unstructured 下载 NLTK 数据），线程无法强杀，
    所以这里只放弃等待——守护线程不会阻止进程退出，也不影响主流程降级。
    """
    result: dict[str, object] = {}

    def run() -> None:
        try:
            result["value"] = func(path)
        except BaseException as exc:  # 线程里的异常要带回主线程
            result["error"] = exc

    worker = threading.Thread(target=run, daemon=True)
    worker.start()
    worker.join(timeout)
    if worker.is_alive():
        raise TimeoutError(f"解析超时（>{timeout:g}s）")
    if "error" in result:
        raise result["error"]  # type: ignore[misc]
    return result["value"]


def parse_resume(
    path: Path,
    engines: tuple[str, ...] = ENGINE_ORDER,
    extractors: dict[str, Extractor] | None = None,
    timeouts: dict[str, float] | None = None,
) -> ParseResult:
    """按优先级尝试各引擎，第一个通过质量闸门的胜出。"""
    available = extractors or EXTRACTORS
    budgets = timeouts or ENGINE_TIMEOUTS
    attempts: list[tuple[str, str]] = []

    for engine in engines:
        try:
            lines, pages = call_with_timeout(
                available[engine], path, budgets.get(engine, DEFAULT_TIMEOUT_S)
            )
            original_pages = pages
        except Exception as exc:  # 兜底：任何解析器异常都算该引擎失败
            message = str(exc).splitlines()[0][:120] if str(exc).strip() else ""
            attempts.append((engine, f"{type(exc).__name__}: {message}".rstrip(": ")))
            continue

        # 先剥掉每页重复的页眉页脚，再用剥离后的内容做质量判定；
        # 原始 pages 传给 assess_quality，让报告里仍能看到剥掉了什么。
        cleaned_pages = strip_repeated_edges(pages)
        if cleaned_pages != pages:
            pages = cleaned_pages
            lines = [line for page in pages for line in page]

        report = assess_quality(lines, original_pages)
        if report.ok:
            return ParseResult(path, engine, lines, pages, report, tuple(attempts))
        attempts.append((engine, "; ".join(report.problems)))

    detail = " | ".join(f"{engine}: {reason}" for engine, reason in attempts)
    raise ParseError(f"所有解析引擎均不可用：{detail}")
