"""解析探针：只回答一个问题——简历能不能被正确读成结构化文本。

不参与主链路，不依赖模型，不产生任何评分。

用法：
    .venv/bin/python probe_parse.py <文件或目录> [...] [--engine native|unstructured|both]

native      = PDF 用 pymupdf、DOCX 用 python-docx
unstructured = 官方解析器（中文/依赖限制下可能失败，失败会如实报告）
both        = 两种都跑，用于对照选型（默认）

输出每个文件的：页数 / 文本层 / 字符数 / 中文占比 / 乱码率 / 章节结构，
并把提取出的纯文本落到 _probe_out/ 供人工核对。
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import pymupdf
from docx import Document

OUT_DIR = Path("_probe_out")
CJK_RE = re.compile(r"[\u4e00-\u9fff]")
# 私用区字符：中文字体缺少 ToUnicode 映射时最常见的乱码形态
PUA_RE = re.compile(r"[\ue000-\uf8ff]")
CTRL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

# 简历常见章节标题，用来判断结构有没有被保住
SECTIONS = (
    "个人信息", "求职意向", "教育", "工作经历", "工作经验", "项目",
    "技能", "专业技能", "实习", "证书", "荣誉", "自我评价", "个人优势",
    "个人总结", "工作描述",
)


def text_stats(text: str) -> dict[str, float | int]:
    """统计一段文本的可用性：字符数、中文占比、乱码率。"""
    total = len(text.strip())
    if total == 0:
        return {"chars": 0, "cjk_ratio": 0.0, "garbled_ratio": 0.0}
    cjk = len(CJK_RE.findall(text))
    garbled = len(PUA_RE.findall(text)) + len(CTRL_RE.findall(text)) + text.count("�")
    return {
        "chars": total,
        "cjk_ratio": round(cjk / total, 3),
        "garbled_ratio": round(garbled / total, 4),
    }


def match_section(line: str) -> str | None:
    """把一行判为章节标题，要求短且命中关键词。"""
    line = line.strip()
    if not line or len(line) > 12:
        return None
    # 带句读的短句是正文换行碎片，不是标题
    if any(mark in line for mark in "。！？，；"):
        return None
    return line if any(key in line for key in SECTIONS) else None


def section_map(lines: list[str]) -> list[tuple[str, int]]:
    """返回 [(章节标题, 该章节下的行数)]，用于判断结构是否被保住。"""
    order: list[str] = []
    counts: dict[str, int] = {}
    current: str | None = None
    for line in lines:
        section = match_section(line)
        if section:
            current = section
            if section not in counts:
                order.append(section)
                counts[section] = 0
        elif current:
            counts[current] += 1
    return [(section, counts[section]) for section in order]


def pdf_page_chars(path: Path) -> tuple[int, list[int], int]:
    """返回 (页数, 每页字符数, 内嵌图片总数)，用于判断是否为扫描件。"""
    with pymupdf.open(path) as doc:
        char_counts = [len(page.get_text("text").strip()) for page in doc]
        images = sum(len(page.get_images(full=True)) for page in doc)
    return len(char_counts), char_counts, images


def extract_native(path: Path) -> list[str]:
    """不依赖第三方解析框架的抽取：pymupdf 读 PDF，python-docx 读 DOCX。"""
    lines: list[str] = []
    if path.suffix.lower() == ".pdf":
        with pymupdf.open(path) as doc:
            for page in doc:
                lines.extend(page.get_text("text").splitlines())
    else:
        document = Document(str(path))
        lines.extend(p.text for p in document.paragraphs)
        # 简历常用表格排版，单元格内容同样要收进来
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    lines.extend(cell.text.splitlines())
    return [line.strip() for line in lines if line.strip()]


def extract_unstructured(path: Path) -> list[str]:
    """官方解析器。中文需要显式指定语言，否则会按英文处理并返回空结果。"""
    from unstructured.partition.auto import partition

    kwargs: dict[str, object] = {"filename": str(path), "languages": ["chi_sim"]}
    if path.suffix.lower() == ".pdf":
        kwargs["strategy"] = "fast"
    elements = partition(**kwargs)
    return [e.text.strip() for e in elements if e.text.strip()]


ENGINES = {"native": extract_native, "unstructured": extract_unstructured}


def probe(path: Path, engine: str) -> None:
    print(f"\n=== {path.name} ===")

    if path.suffix.lower() == ".pdf":
        pages, char_counts, images = pdf_page_chars(path)
        print(f"页数 {pages} | 每页字符 {char_counts} | 内嵌图片 {images}")
        if pages and max(char_counts) < 50:
            print("  !! 文本层几乎为空，疑似扫描件/图片型 PDF，需要 OCR")

    for name in ("native", "unstructured") if engine == "both" else (engine,):
        try:
            lines = ENGINES[name](path)
        except Exception as exc:  # 解析失败必须显式失败，不能猜
            print(f"  [{name}] 失败: {type(exc).__name__}: {str(exc).splitlines()[0][:120]}")
            continue
        text = "\n".join(lines)
        stats = text_stats(text)
        print(
            f"  [{name}] {stats['chars']} 字符 | 中文 {stats['cjk_ratio']:.1%} "
            f"| 乱码 {stats['garbled_ratio']:.2%} | 行 {len(lines)}"
        )
        print(f"  [{name}] 章节: {section_map(lines) or '未识别到'}")
        OUT_DIR.mkdir(exist_ok=True)
        (OUT_DIR / f"{path.stem}.{name}.txt").write_text(text, encoding="utf-8")


def collect(targets: list[str]) -> list[Path]:
    files: list[Path] = []
    for raw in targets:
        path = Path(raw).expanduser()
        if path.is_dir():
            files.extend(sorted(p for p in path.iterdir()
                                if p.suffix.lower() in {".pdf", ".docx"}))
        elif path.is_file():
            files.append(path)
        else:
            print(f"跳过不存在的路径: {path}")
    return files


def main(argv: list[str]) -> int:
    engine = "both"
    if "--engine" in argv:
        index = argv.index("--engine")
        engine = argv[index + 1]
        del argv[index:index + 2]
    if engine not in {"native", "unstructured", "both"}:
        print(f"未知 engine: {engine}")
        return 2

    files = collect(argv)
    if not files:
        print("没有可解析的文件")
        return 1
    for path in files:
        probe(path, engine)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
